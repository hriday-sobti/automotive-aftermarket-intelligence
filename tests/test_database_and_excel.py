# Unit tests for relational database tables, SQL analytical views, and Excel scenario models.

import os
import sqlite3
import pytest
import openpyxl

from src.analytics.excel_scenario_generator import generate_excel_scenario_model
from src.utils.config_utils import load_config


@pytest.fixture
def db_conn():
    config = load_config()
    db_path = config["database"]["sqlite_path"]
    conn = sqlite3.connect(db_path)
    yield conn
    conn.close()


def test_database_table_existence(db_conn):
    cur = db_conn.cursor()
    tables = [
        "dim_date", "dim_region", "dim_vehicle_segment", "dim_product",
        "dim_distributor", "dim_workshop", "dim_promotion", "dim_sku_pareto",
        "fact_sales", "fact_inventory", "fact_demand"
    ]
    for t in tables:
        res = cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        assert res > 0, f"Table {t} should contain rows"


def test_sql_analytical_views_existence(db_conn):
    cur = db_conn.cursor()
    views = [
        "v_monthly_sales_performance",
        "v_category_regional_mix",
        "v_distributor_performance_ranking",
        "v_trade_promotion_evaluation",
        "v_inventory_stockout_summary"
    ]
    for v in views:
        res = cur.execute(f"SELECT COUNT(*) FROM {v}").fetchone()[0]
        assert res > 0, f"View {v} should return rows"


def test_sql_zero_orphan_skus(db_conn):
    cur = db_conn.cursor()
    q = """
    SELECT COUNT(*) FROM fact_sales s
    LEFT JOIN dim_product p ON s.sku_id = p.sku_id
    WHERE p.sku_id IS NULL;
    """
    assert cur.execute(q).fetchone()[0] == 0


def test_sql_revenue_reconciliation(db_conn):
    cur = db_conn.cursor()
    q = "SELECT COUNT(*) FROM fact_sales WHERE ABS(revenue - ROUND(quantity * net_unit_price, 2)) > 0.05;"
    assert cur.execute(q).fetchone()[0] == 0


def test_sql_gross_profit_reconciliation(db_conn):
    cur = db_conn.cursor()
    q = "SELECT COUNT(*) FROM fact_sales WHERE ABS(gross_profit - ROUND(revenue - cost, 2)) > 0.05;"
    assert cur.execute(q).fetchone()[0] == 0


def test_sql_inventory_balance_reconciliation(db_conn):
    cur = db_conn.cursor()
    q = "SELECT COUNT(*) FROM fact_inventory WHERE closing_stock != (opening_stock + receipts - units_sold);"
    assert cur.execute(q).fetchone()[0] == 0


def test_sql_no_negative_closing_stock(db_conn):
    cur = db_conn.cursor()
    q = "SELECT COUNT(*) FROM fact_inventory WHERE closing_stock < 0;"
    assert cur.execute(q).fetchone()[0] == 0


def test_excel_scenario_model_generation(tmp_path):
    output_xlsx = str(tmp_path / "test_model.xlsx")
    generate_excel_scenario_model(output_xlsx)
    assert os.path.exists(output_xlsx)

    wb = openpyxl.load_workbook(output_xlsx, data_only=False)
    assert "Promotion Scenario Model" in wb.sheetnames
    ws = wb["Promotion Scenario Model"]

    # Verify input values
    assert ws["C6"].value == 15000  # Baseline units
    assert ws["C7"].value == 1850.0 # List price
    assert ws["C8"].value == 1150.0 # Cost
    assert ws["C9"].value == 0.05   # Base discount

    # Verify formulas
    assert "=ROUND(C6*(1+C14), 0)" in ws["H7"].value
    assert "=H9-H10" in ws["H11"].value
    assert "=I11/H14" in ws["C19"].value

    # Verify sensitivity matrix formulas
    assert ws["G19"].value.startswith("=ROUND(")


@pytest.mark.parametrize("cell,expected_val", [
    ("B5", "1. BASELINE ASSUMPTIONS (NON-PROMOTED)"),
    ("B12", "2. PROMOTIONAL PROPOSAL INPUTS"),
    ("F5", "3. BASELINE VS. PROMOTION FINANCIAL RECONCILIATION"),
    ("B17", "4. EXECUTIVE DECISION EVALUATION"),
    ("F17", "5. SENSITIVITY TABLE: NET INCREMENTAL GROSS PROFIT (₹)"),
])
def test_excel_headers(tmp_path, cell, expected_val):
    output_xlsx = str(tmp_path / "test_headers.xlsx")
    generate_excel_scenario_model(output_xlsx)
    wb = openpyxl.load_workbook(output_xlsx, data_only=False)
    ws = wb["Promotion Scenario Model"]
    assert ws[cell].value == expected_val
