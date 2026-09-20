"""Excel Commercial Scenario Model Generator.

Builds a fully functioning, professional Microsoft Excel workbook (.xlsx)
with dynamic spreadsheet formulas, explicit input parameters, calculation blocks,
and scenario sensitivity tables for aftermarket trade promotion planning.

Uses openpyxl with clean financial formatting, bold headers, and visible color hierarchy:
- INPUTS: Pastel Yellow / Cream background
- CALCULATIONS: Light Slate / Gray background
- OUTPUTS / KPIS: Soft Mint Green background
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "../..")))


def generate_excel_scenario_model(output_path: str = "excel/scenario_model/trade_promotion_scenario_model.xlsx"):
    """Generate professional Excel trade promotion sensitivity model."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Promotion Scenario Model"

    # Style definitions
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E79")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="595959")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_regular = Font(name="Segoe UI", size=10)

    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_input = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft yellow
    fill_output = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft green
    fill_calc = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Light gray

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1F4E79"))

    # Title Block
    ws["B2"] = "Automotive Aftermarket Commercial Promotion Scenario Simulator"
    ws["B2"].font = font_title
    ws["B3"] = "Interactive B2B trade discount sensitivity & net incremental gross profit model"
    ws["B3"].font = font_subtitle

    # Section 1: Baseline Inputs (Cells B5:D10)
    ws["B5"] = "1. BASELINE ASSUMPTIONS (NON-PROMOTED)"
    ws["B5"].font = font_bold
    ws["B5"].fill = fill_navy
    ws["B5"].font = font_header
    ws.merge_cells("B5:D5")

    inputs_baseline = [
        ("B6", "Baseline Period Units (Run-rate)", "C6", 15000, "#,##0"),
        ("B7", "Product List Price (₹)", "C7", 1850.00, "₹#,##0.00"),
        ("B8", "Unit Cost of Goods Sold (₹)", "C8", 1150.00, "₹#,##0.00"),
        ("B9", "Standard Wholesale Trade Discount %", "C9", 0.05, "0.0%"),
        ("B10", "Campaign Fixed Promotional Budget (₹)", "C10", 350000.00, "₹#,##0.00")
    ]

    for label_cell, label, val_cell, val, fmt in inputs_baseline:
        ws[label_cell] = label
        ws[label_cell].font = font_regular
        ws[val_cell] = val
        ws[val_cell].font = font_bold
        ws[val_cell].fill = fill_input
        ws[val_cell].number_format = fmt
        ws[val_cell].alignment = Alignment(horizontal="right")
        ws[label_cell].border = border_cell
        ws[val_cell].border = border_cell

    # Section 2: Promotional Proposal Inputs (Cells B12:D14)
    ws["B12"] = "2. PROMOTIONAL PROPOSAL INPUTS"
    ws["B12"].font = font_header
    ws["B12"].fill = fill_navy
    ws.merge_cells("B12:D12")

    inputs_promo = [
        ("B13", "Proposed Promotional Discount %", "C13", 0.12, "0.0%"),
        ("B14", "Expected Observed Volume Uplift %", "C14", 0.28, "0.0%")
    ]

    for label_cell, label, val_cell, val, fmt in inputs_promo:
        ws[label_cell] = label
        ws[label_cell].font = font_regular
        ws[val_cell] = val
        ws[val_cell].font = font_bold
        ws[val_cell].fill = fill_input
        ws[val_cell].number_format = fmt
        ws[val_cell].alignment = Alignment(horizontal="right")
        ws[label_cell].border = border_cell
        ws[val_cell].border = border_cell

    # Section 3: Dynamic Model Calculations (Cells F5:I14)
    ws["F5"] = "3. BASELINE VS. PROMOTION FINANCIAL RECONCILIATION"
    ws["F5"].font = font_header
    ws["F5"].fill = fill_navy
    ws.merge_cells("F5:I5")

    headers = [("F6", "Financial Metric"), ("G6", "Baseline Model"), ("H6", "Promoted Model"), ("I6", "Incremental Variance")]
    for h_cell, h_text in headers:
        ws[h_cell] = h_text
        ws[h_cell].font = font_bold
        ws[h_cell].fill = fill_calc
        ws[h_cell].border = border_cell

    calc_rows = [
        ("Total Units Sold", "=C6", "=ROUND(C6*(1+C14), 0)", "=H7-G7", "#,##0"),
        ("Net Selling Price / Unit", "=C7*(1-C9)", "=C7*(1-C13)", "=H8-G8", "₹#,##0.00"),
        ("Gross Revenue", "=G7*G8", "=H7*H8", "=H9-G9", "₹#,##0.00"),
        ("Total Cost of Goods Sold", "=G7*C8", "=H7*C8", "=H10-G10", "₹#,##0.00"),
        ("Gross Profit", "=G9-G10", "=H9-H10", "=H11-G11", "₹#,##0.00"),
        ("Realized Gross Margin %", "=G11/G9", "=H11/H9", "=H12-G12", "0.0%"),
        ("Trade Discount Concession Cost", "=G7*(C7-G8)", "=H7*(C7-H8)", "=H13-G13", "₹#,##0.00"),
        ("Total Promotion Cost (Discount + Budget)", 0.00, "=H13+C10", "=H14-G14", "₹#,##0.00"),
    ]

    for idx, (label, f_base, f_promo, f_diff, fmt) in enumerate(calc_rows, start=7):
        ws[f"F{idx}"] = label
        ws[f"F{idx}"].font = font_regular
        ws[f"F{idx}"].border = border_cell

        for col, formula in [("G", f_base), ("H", f_promo), ("I", f_diff)]:
            c = f"{col}{idx}"
            ws[c] = formula
            ws[c].font = font_bold if "Profit" in label else font_regular
            ws[c].number_format = fmt
            ws[c].border = border_cell
            ws[c].alignment = Alignment(horizontal="right")
            if "Profit" in label:
                ws[c].fill = fill_output

    # Section 4: Key Decision Indicators (Cells B17:D21)
    ws["B17"] = "4. EXECUTIVE DECISION EVALUATION"
    ws["B17"].font = font_header
    ws["B17"].fill = fill_navy
    ws.merge_cells("B17:D17")

    kpi_rows = [
        ("B18", "Net Incremental Gross Profit (₹)", "C18", "=I11", "₹#,##0.00"),
        ("B19", "Promotion Effectiveness Index (PEI)", "C19", "=I11/H14", "0.00x"),
        ("B20", "Gross Margin Dilution (Basis Points)", "C20", "=(G12-H12)*10000", "#,##0 bps"),
        ("B21", "Commercial Recommendation", "C21", '=IF(C19>1.0, "APPROVE: Value Accretive", IF(C19>0.0, "REVIEW: Volume Driver / Margin Dilutive", "REJECT: Value Destructive"))', "@")
    ]

    for label_cell, label, val_cell, formula, fmt in kpi_rows:
        ws[label_cell] = label
        ws[label_cell].font = font_bold
        ws[val_cell] = formula
        ws[val_cell].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
        ws[val_cell].fill = fill_output
        ws[val_cell].number_format = fmt
        ws[val_cell].alignment = Alignment(horizontal="right")
        ws[label_cell].border = border_cell
        ws[val_cell].border = border_cell

    # Section 5: Sensitivity Data Table (Discount % vs Volume Uplift %)
    ws["F17"] = "5. SENSITIVITY TABLE: NET INCREMENTAL GROSS PROFIT (₹)"
    ws["F17"].font = font_header
    ws["F17"].fill = fill_navy
    ws.merge_cells("F17:L17")

    discounts = [0.08, 0.10, 0.12, 0.14, 0.16]
    uplifts = [0.15, 0.20, 0.25, 0.30, 0.35, 0.40]

    ws["F18"] = "Discount \\ Uplift"
    ws["F18"].font = font_bold
    ws["F18"].fill = fill_calc
    ws["F18"].border = border_cell

    for c_idx, up in enumerate(uplifts, start=7):
        col_let = get_column_letter(c_idx)
        ws[f"{col_let}18"] = up
        ws[f"{col_let}18"].number_format = "0.0%"
        ws[f"{col_let}18"].font = font_bold
        ws[f"{col_let}18"].fill = fill_calc
        ws[f"{col_let}18"].alignment = Alignment(horizontal="center")
        ws[f"{col_let}18"].border = border_cell

    for r_idx, disc in enumerate(discounts, start=19):
        ws[f"F{r_idx}"] = disc
        ws[f"F{r_idx}"].number_format = "0.0%"
        ws[f"F{r_idx}"].font = font_bold
        ws[f"F{r_idx}"].fill = fill_calc
        ws[f"F{r_idx}"].alignment = Alignment(horizontal="center")
        ws[f"F{r_idx}"].border = border_cell

        for c_idx, up in enumerate(uplifts, start=7):
            col_let = get_column_letter(c_idx)
            # Dynamic spreadsheet formula for each matrix coordinate:
            # PromoUnits = C6*(1+Up)
            # NetPrice = C7*(1-Disc)
            # PromoGP = PromoUnits * (NetPrice - C8)
            # BaseGP = G11
            # IncGP = PromoGP - BaseGP
            cell_formula = f"=ROUND(($C$6*(1+{col_let}$18))*(($C$7*(1-$F{r_idx}))-$C$8)-$G$11, 0)"
            ws[f"{col_let}{r_idx}"] = cell_formula
            ws[f"{col_let}{r_idx}"].number_format = "₹#,##0"
            ws[f"{col_let}{r_idx}"].font = font_regular
            ws[f"{col_let}{r_idx}"].alignment = Alignment(horizontal="right")
            ws[f"{col_let}{r_idx}"].border = border_cell

    # Adjust Column Widths
    col_widths = {
        "A": 3, "B": 38, "C": 18, "D": 4, "E": 4,
        "F": 35, "G": 20, "H": 22, "I": 22, "J": 16, "K": 16, "L": 16
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print(f"Scenario model saved to {output_path}")


if __name__ == "__main__":
    generate_excel_scenario_model()
